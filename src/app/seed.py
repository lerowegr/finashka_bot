"""Seed-данные: то, чего нет на tabiturient.ru, переносится из старых Excel.

Это авторский контент прежней версии бота: описания направлений, «кем работать»,
военные кафедры и текстовые описания вузов.
"""

import json
import logging
import re
import sqlite3
from pathlib import Path

from sqlalchemy import select

from app.config import SEED_DIR
from app.db import models as m
from app.db.engine import session_factory

log = logging.getLogger(__name__)

_CODE_RE = re.compile(r"^\d{2}\.\d{2}\.\d{2}$")
_UGSN_RE = re.compile(r"^\d{2}\.00\.00$")
_PSEUDO_SECTIONS = {
    "Специалитет",
    "Специалитет, содержащий сведения гос.тайны или служебную информацию",
}


def extract_from_excel(old_dir: Path) -> None:
    """Одноразовое извлечение из finashka_bot-main/{spec,vuz}.xlsx в data/seed/*.json."""
    import openpyxl

    SEED_DIR.mkdir(parents=True, exist_ok=True)

    directions: dict[str, dict] = {}
    sections: dict[str, list[str]] = {}
    current_section = None

    wb = openpyxl.load_workbook(old_dir / "spec.xlsx", read_only=True)
    sheet = wb[wb.sheetnames[1]]  # «Описание специальностей»
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not row[0]:
            continue
        first = str(row[0]).strip()
        if _UGSN_RE.match(first):
            if current_section:
                sections.setdefault(current_section, []).append(first)
        elif _CODE_RE.match(first):
            directions[first] = {
                "name": str(row[1] or "").strip(),
                "description": str(row[2] or "").strip() or None,
                "career": str(row[3] or "").strip() or None,
            }
        elif not row[1] and first not in _PSEUDO_SECTIONS:
            current_section = first
    wb.close()

    universities: dict[str, dict] = {}
    wb = openpyxl.load_workbook(old_dir / "vuz.xlsx", read_only=True)
    for i, row in enumerate(wb[wb.sheetnames[0]].iter_rows(values_only=True)):
        if i == 0 or not row[0]:
            continue
        universities[str(row[0]).strip()] = {
            "description": str(row[2] or "").strip() or None,
            "military": str(row[11] or "").strip().lower() == "да",
        }
    wb.close()

    for name, payload in (
        ("directions.json", directions),
        ("sections.json", sections),
        ("universities.json", universities),
    ):
        (SEED_DIR / name).write_text(
            json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8"
        )
    log.info(
        "seed извлечён: %d направлений, %d разделов, %d вузов",
        len(directions),
        len(sections),
        len(universities),
    )


def _norm_name(s: str) -> str:
    s = s.lower().replace("ё", "е")
    s = re.sub(r"[«»\"']", "", s)
    return re.sub(r"\s+", " ", s).strip()


async def apply_seed() -> None:
    """Накладывает seed на БД, уже наполненную ETL-ом."""
    directions = json.loads((SEED_DIR / "directions.json").read_text(encoding="utf-8"))
    sections = json.loads((SEED_DIR / "sections.json").read_text(encoding="utf-8"))
    universities = json.loads((SEED_DIR / "universities.json").read_text(encoding="utf-8"))

    async with session_factory() as session:
        matched_dir = 0
        for code, payload in directions.items():
            direction = (
                await session.execute(select(m.Direction).where(m.Direction.code == code))
            ).scalar_one_or_none()
            if direction is not None:
                direction.description = payload["description"]
                direction.career = payload["career"]
                matched_dir += 1

        for section_name, ugsn_codes in sections.items():
            section = (
                await session.execute(select(m.Section).where(m.Section.name == section_name))
            ).scalar_one_or_none()
            if section is None:
                section = m.Section(name=section_name)
                session.add(section)
                await session.flush()
            for code in ugsn_codes:
                ugsn = await session.get(m.Ugsn, code)
                if ugsn is not None:
                    ugsn.section_id = section.id

        all_vuz = (await session.execute(select(m.University))).scalars().all()
        by_norm = {_norm_name(v.full_name): v for v in all_vuz}
        matched_vuz = 0
        unmatched: list[str] = []
        for name, payload in universities.items():
            key = _norm_name(name)
            vuz = by_norm.get(key)
            if vuz is None:  # частичное совпадение
                candidates = [v for k, v in by_norm.items() if key in k or k in key]
                vuz = candidates[0] if len(candidates) == 1 else None
            if vuz is None:
                unmatched.append(name)
                continue
            vuz.military_center = payload["military"]
            if payload["description"]:
                vuz.description = payload["description"]
            matched_vuz += 1

        await session.commit()

    log.info(
        "seed применён: %d/%d направлений, %d/%d вузов",
        matched_dir,
        len(directions),
        matched_vuz,
        len(universities),
    )
    if unmatched:
        log.warning("не сматчились вузы (%d): %s", len(unmatched), "; ".join(unmatched[:10]))


async def migrate_favorites(old_db: Path) -> None:
    """Перенос избранного из старой SQLite (таблица fav: id, code)."""
    conn = sqlite3.connect(old_db)
    rows = conn.execute("SELECT id, code FROM fav").fetchall()
    conn.close()

    moved = skipped = 0
    async with session_factory() as session:
        for tg_id, code in rows:
            direction = (
                await session.execute(select(m.Direction).where(m.Direction.code == code))
            ).scalar_one_or_none()
            if direction is None:
                skipped += 1
                continue
            if await session.get(m.User, tg_id) is None:
                session.add(m.User(tg_id=tg_id))
                await session.flush()
            if await session.get(m.Favorite, (tg_id, "direction", direction.id)) is None:
                session.add(m.Favorite(tg_id=tg_id, kind="direction", item_id=direction.id))
                moved += 1
        await session.commit()
    log.info("избранное перенесено: %d записей, пропущено %d", moved, skipped)
